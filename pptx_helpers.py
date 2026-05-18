from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

NAVY  = RGBColor(0x0F,0x3C,0x60); DARK  = RGBColor(0x1A,0x1A,0x2E)
RED   = RGBColor(0xE9,0x45,0x60); GOLD  = RGBColor(0xF5,0xA6,0x23)
WHITE = RGBColor(0xFF,0xFF,0xFF); LGRAY = RGBColor(0xF4,0xF6,0xFB)
MGRAY = RGBColor(0xCC,0xD3,0xE8); TEAL  = RGBColor(0x00,0xB4,0xD8)
GREEN = RGBColor(0x2D,0xC6,0x5E); ORANGE= RGBColor(0xFF,0x7B,0x25)

PHASE_COLORS = {1:TEAL, 2:GREEN, 3:ORANGE, 4:RED, 5:GOLD}
DAY_COLORS   = {0:NAVY,1:TEAL,2:GREEN,3:ORANGE,4:RED,5:GOLD}

def new_prs():
    p=Presentation(); p.slide_width=Inches(13.33); p.slide_height=Inches(7.5); return p

def blank(prs): return prs.slides.add_slide(prs.slide_layouts[6])

def rect(s,l,t,w,h,fill=None,line=None):
    sh=s.shapes.add_shape(1,Inches(l),Inches(t),Inches(w),Inches(h))
    sh.line.fill.background()
    if fill: sh.fill.solid(); sh.fill.fore_color.rgb=fill
    else: sh.fill.background()
    if line: sh.line.color.rgb=line; sh.line.width=Pt(1.5)
    return sh

def txt(s,text,l,t,w,h,size=Pt(12),bold=False,color=DARK,align=PP_ALIGN.LEFT,italic=False):
    tb=s.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h))
    tb.word_wrap=True; tf=tb.text_frame; tf.word_wrap=True
    p=tf.paragraphs[0]; p.alignment=align
    r=p.add_run(); r.text=text; r.font.size=size
    r.font.bold=bold; r.font.color.rgb=color; r.font.italic=italic
    return tb

def pill(s,text,l,t,w,h,bg=NAVY,fg=WHITE,size=Pt(10)):
    rect(s,l,t,w,h,fill=bg)
    txt(s,text,l,t,w,h,size=size,bold=True,color=fg,align=PP_ALIGN.CENTER)

def header_bar(s,title,accent=NAVY):
    rect(s,0,0,13.33,1.0,fill=DARK)
    rect(s,0,0,0.18,7.5,fill=accent)
    txt(s,title,0.4,0.15,12.5,0.72,size=Pt(24),bold=True,color=WHITE)

def footer(s,day_num,total=78,label="TAMOU NEURAL PATH"):
    rect(s,0,7.1,13.33,0.4,fill=DARK)
    txt(s,f"{label}  |  Jour {day_num} / {total}",0.3,7.12,10,0.3,size=Pt(9),color=MGRAY)
    txt(s,f"{day_num}/{total}",12.2,7.12,1.0,0.3,size=Pt(9),color=MGRAY,align=PP_ALIGN.RIGHT)

def bullet_rows(s,items,l,t,w,accent=NAVY,row_h=0.72):
    for i,item in enumerate(items):
        y=t+i*row_h
        bg=LGRAY if i%2==0 else WHITE
        rect(s,l,y,w,row_h-0.06,fill=bg)
        rect(s,l,y,0.1,row_h-0.06,fill=accent)
        if isinstance(item,tuple):
            title,body=item
            txt(s,title,l+0.22,y+0.05,w-0.3,0.3,size=Pt(10),bold=True,color=accent)
            txt(s,body, l+0.22,y+0.36,w-0.3,0.32,size=Pt(9.5),color=DARK)
        else:
            txt(s,"•  "+item,l+0.22,y+0.18,w-0.3,0.42,size=Pt(10),color=DARK)

def step_cards(s,steps,l=0.35,t=1.15,cols=2):
    n=len(steps); col_w=(12.6-(cols-1)*0.15)/cols
    for i,(icon,title,body,col) in enumerate(steps):
        col_i=i%cols; row_i=i//cols
        xl=l+col_i*(col_w+0.15); y=t+row_i*1.65
        rect(s,xl,y,col_w,1.52,fill=LGRAY)
        rect(s,xl,y,0.12,1.52,fill=col)
        txt(s,icon,xl+0.18,y+0.12,0.55,0.55,size=Pt(20))
        txt(s,title,xl+0.78,y+0.14,col_w-0.95,0.38,size=Pt(11),bold=True,color=col)
        txt(s,body, xl+0.18,y+0.62,col_w-0.3,0.82,size=Pt(9.5),color=DARK)

def prompt_box(s,prompt_text,l=0.35,t=1.1,w=12.6,h=2.2):
    rect(s,l,t,w,h,fill=DARK)
    rect(s,l,t,0.15,h,fill=RED)
    lines=prompt_text.split('\n')
    for i,line in enumerate(lines):
        txt(s,line,l+0.28,t+0.18+i*0.32,w-0.45,0.3,
            size=Pt(10),color=RGBColor(0xA8,0xF0,0xE0))

def make_day(data):
    """
    data = dict with keys:
      day_num, week, phase, phase_name, day_name,
      title, subtitle, accent_color,
      theory_title, theory_points (list of str or (title,body)),
      theory_insight,
      exercise_title, exercise_steps (list of (icon,title,body,color)),
      production_title, production_items (list),
      livrable, livrable_desc,
      tomorrow_title, tomorrow_desc
    """
    prs=new_prs(); ac=data.get('accent_color',TEAL)
    ph=data['phase']; d=data['day_num']

    # ── SLIDE 1 COVER ──────────────────────────────────────
    s=blank(prs)
    rect(s,0,0,13.33,7.5,fill=DARK)
    rect(s,0,0,0.18,7.5,fill=ac)
    rect(s,0,6.6,13.33,0.9,fill=NAVY)
    txt(s,"TAMOU NEURAL PATH",1,0.9,11,0.8,size=Pt(32),bold=True,color=WHITE,align=PP_ALIGN.CENTER)
    rect(s,4,1.85,5.3,0.07,fill=ac)
    txt(s,f"JOUR {d}  —  {data['day_name'].upper()}",1,2.05,11,0.72,
        size=Pt(26),bold=True,color=GOLD,align=PP_ALIGN.CENTER)
    txt(s,data['title'],1,2.9,11,0.55,size=Pt(17),color=WHITE,align=PP_ALIGN.CENTER)
    txt(s,data.get('subtitle',''),1,3.52,11,0.42,size=Pt(12),italic=True,color=MGRAY,align=PP_ALIGN.CENTER)
    for i,(ic,lb) in enumerate([(f"Phase {ph}",data['phase_name']),
                                 (f"Sem. {data['week']}",""),
                                 (f"J{d}/78","")]):
        xl=2.2+i*3.1
        rect(s,xl,4.6,2.5,0.85,fill=NAVY)
        txt(s,ic,xl+0.15,4.68,2.2,0.38,size=Pt(13),bold=True,color=ac)
        if lb: txt(s,lb,xl+0.15,5.05,2.2,0.32,size=Pt(9),color=MGRAY)
    txt(s,f"1/7",12.5,7.15,0.8,0.3,size=Pt(9),color=MGRAY,align=PP_ALIGN.RIGHT)

    # ── SLIDE 2 PLANNING ──────────────────────────────────
    s=blank(prs)
    header_bar(s,"PLANNING DU JOUR",ac)
    rect(s,0,0,0.18,7.5,fill=ac)
    blocks=[
        (TEAL, "🌅 MATIN",     "09h00–09h45","THÉORIE",  data['theory_title']),
        (TEAL, "",             "09h45–11h00","EXERCICE",  data['exercise_title']),
        (ORANGE,"🌇 APRÈS-MIDI","14h00–16h00","PRODUCTION",data['production_title']),
        (GREEN, "📋 LIVRABLE", "16h00",      "LIVRABLE",  data['livrable']),
    ]
    for i,(col,label,heure,badge,desc) in enumerate(blocks):
        y=1.18+i*1.42
        rect(s,0.35,y,12.6,1.3,fill=LGRAY)
        rect(s,0.35,y,0.12,1.3,fill=col)
        if label: txt(s,label,0.6,y+0.1,2.2,0.42,size=Pt(11),bold=True,color=col)
        txt(s,heure,0.6,y+0.7,2.2,0.38,size=Pt(10),color=MGRAY)
        pill(s,badge,3.0,y+0.4,1.7,0.44,bg=col)
        txt(s,desc,4.9,y+0.38,7.8,0.5,size=Pt(11),color=DARK)
    footer(s,d)

    # ── SLIDE 3 THÉORIE ───────────────────────────────────
    s=blank(prs)
    header_bar(s,f"THÉORIE — {data['theory_title']}",ac)
    rect(s,0,0,0.18,7.5,fill=ac)
    pill(s,"THÉORIE  45 min",0.35,1.08,2.5,0.38,bg=ac)
    bullet_rows(s,data['theory_points'],l=0.35,t=1.55,w=12.6,accent=ac)
    # insight
    ins_y=1.55+len(data['theory_points'])*0.72+0.1
    if ins_y<6.8:
        rect(s,0.35,ins_y,12.6,0.62,fill=NAVY)
        txt(s,"💡  "+data['theory_insight'],0.5,ins_y+0.1,12.2,0.42,
            size=Pt(11),bold=True,color=GOLD)
    footer(s,d)

    # ── SLIDE 4 EXERCICE ──────────────────────────────────
    s=blank(prs)
    header_bar(s,f"EXERCICE — {data['exercise_title']}",ac)
    rect(s,0,0,0.18,7.5,fill=ac)
    pill(s,"EXERCICE  75 min",0.35,1.08,2.5,0.38,bg=TEAL)
    step_cards(s,data['exercise_steps'],l=0.35,t=1.55)
    footer(s,d)

    # ── SLIDE 5 PRODUCTION ────────────────────────────────
    s=blank(prs)
    header_bar(s,f"PRODUCTION — {data['production_title']}",ac)
    rect(s,0,0,0.18,7.5,fill=ac)
    pill(s,"PRODUCTION  2h",0.35,1.08,2.5,0.38,bg=ORANGE)
    bullet_rows(s,data['production_items'],l=0.35,t=1.55,w=12.6,accent=ORANGE)
    footer(s,d)

    # ── SLIDE 6 LIVRABLE ──────────────────────────────────
    s=blank(prs)
    rect(s,0,0,13.33,7.5,fill=LGRAY)
    rect(s,0,0,0.18,7.5,fill=GREEN)
    rect(s,0,0,13.33,1.0,fill=DARK)
    txt(s,f"✅  LIVRABLE DU JOUR {d}",0.4,0.15,12.5,0.72,size=Pt(24),bold=True,color=WHITE)
    rect(s,0.35,1.1,12.6,1.5,fill=GREEN)
    txt(s,data['livrable'],0.55,1.2,12.2,0.55,size=Pt(18),bold=True,color=WHITE)
    txt(s,data['livrable_desc'],0.55,1.78,12.2,0.7,size=Pt(11),color=WHITE,italic=True)
    txt(s,"📌  Sauvegardez dans Notion → Portfolio → Phase "+str(ph),
        0.55,2.72,12.2,0.42,size=Pt(11),bold=True,color=NAVY)
    # checklist
    checks=["Livrable sauvegardé dans Notion","Journal de bord rempli","Meilleur prompt archivé","Niveau de confiance noté /10"]
    for i,c in enumerate(checks):
        rect(s,0.35,3.3+i*0.72,12.6,0.66,fill=WHITE if i%2 else LGRAY)
        txt(s,"⬜  "+c,0.55,3.38+i*0.72,12.0,0.48,size=Pt(11),color=DARK)
    footer(s,d)

    # ── SLIDE 7 DEMAIN ────────────────────────────────────
    s=blank(prs)
    rect(s,0,0,13.33,7.5,fill=DARK)
    rect(s,0,0,0.18,7.5,fill=ac)
    txt(s,f"DEMAIN — JOUR {d+1}",0.5,1.2,12,0.7,size=Pt(28),bold=True,color=WHITE,align=PP_ALIGN.CENTER)
    txt(s,data['tomorrow_title'],0.5,2.05,12,0.7,size=Pt(22),bold=True,color=GOLD,align=PP_ALIGN.CENTER)
    rect(s,3.5,2.88,6.3,0.07,fill=ac)
    txt(s,data['tomorrow_desc'],1,3.1,11,1.0,size=Pt(14),color=MGRAY,align=PP_ALIGN.CENTER,italic=True)
    rect(s,0.35,4.5,12.6,1.4,fill=NAVY)
    txt(s,"📋  RAPPELS AVANT DE FERMER :",0.55,4.6,12,0.4,size=Pt(12),bold=True,color=GOLD)
    txt(s,"→  Journal de bord Notion rempli   →   Livrable sauvegardé   →   Niveau confiance /10 noté",
        0.55,5.05,12,0.42,size=Pt(11),color=WHITE)
    txt(s,f"TAMOU NEURAL PATH  |  Jour {d} sur 78  |  Phase {ph} — {data['phase_name']}",
        0.5,7.12,12,0.3,size=Pt(9),color=MGRAY,align=PP_ALIGN.CENTER)

    return prs
